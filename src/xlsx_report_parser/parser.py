"""Parser for wide XLSX reports: one row per device, one column per day."""

from __future__ import annotations

import math
from collections.abc import Iterable
from datetime import date
from os import PathLike
from typing import IO

from openpyxl import load_workbook

from .models import DeviceInfo, DeviceRow, ParseWarning, Report, ReportFormatError
from .schema import METRICS, MetricSpec, field_for_header, normalise_unit, parse_date_header

Source = str | PathLike | IO[bytes]


def _is_blank(cell: object) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _to_text(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))  # "1234" instead of "1234.0" for ids stored as numbers
    return str(cell).strip()


def coerce_number(cell: object) -> tuple[float | None, str | None]:
    """Convert a cell into a float.

    Returns ``(value, problem)``. Empty cells and dash placeholders yield ``(None, None)``;
    unusable content yields ``(None, "reason")`` so the caller can record a warning.
    """
    if _is_blank(cell):
        return None, None
    if isinstance(cell, bool):
        return None, "boolean is not a number"
    if isinstance(cell, int | float):
        number = float(cell)
    else:
        text = str(cell).strip().replace("\xa0", "").replace(" ", "")
        if text in {"-", "—", "–", "n/a", "N/A"}:
            return None, None
        try:
            number = float(text.replace(",", "."))
        except ValueError:
            return None, f"not a number: {str(cell)[:30]!r}"
    if math.isnan(number) or math.isinf(number):
        return None, "NaN/inf is not allowed"
    if number < 0:
        return None, f"negative value {number:g}"
    return number, None


def _read_rows(source: Source, sheet: str | None) -> tuple[str, Iterable[tuple]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
    except KeyError as exc:
        names = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ReportFormatError(f"sheet {sheet!r} not found (available: {names})") from exc
    # materialise so the workbook can be closed (read-only mode keeps the file handle open)
    rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    title = worksheet.title
    workbook.close()
    return title, rows


def _select_date_columns(
    header: tuple, spec: MetricSpec, energy_unit: str | None, warnings: list[ParseWarning]
) -> dict[date, tuple[int, float, str]]:
    """Pick exactly one column per date: ``{day: (column_index, factor, label)}``."""
    default_unit = normalise_unit(energy_unit) if energy_unit else spec.default_unit
    chosen: dict[date, tuple[int, float, str, bool]] = {}

    for index, cell in enumerate(header):
        parsed = parse_date_header(cell)
        if parsed is None:
            continue
        day, unit = parsed
        explicit = unit is not None
        unit = unit or default_unit
        if unit not in spec.factors:
            warnings.append(
                ParseWarning(None, f"ignored column with unit {unit!r}", column=str(cell))
            )
            continue
        candidate = (index, spec.factors[unit], str(cell), explicit and spec.factors[unit] == 1.0)
        previous = chosen.get(day)
        if previous is None:
            chosen[day] = candidate
        elif previous[3] == candidate[3] and previous[1] == candidate[1]:
            raise ReportFormatError(f"duplicate column for {day.isoformat()} ({cell!r})")
        elif candidate[3] and not previous[3]:
            chosen[day] = candidate  # prefer the column already expressed in the target unit

    return {day: (idx, factor, label) for day, (idx, factor, label, _) in chosen.items()}


def parse_report(
    source: Source,
    kind: str,
    *,
    sheet: str | None = None,
    energy_unit: str | None = None,
) -> Report:
    """Parse one report.

    :param source: path or binary file object of an ``.xlsx`` workbook
    :param kind: ``"energy"``, ``"uptime"`` or ``"hashrate"``
    :param sheet: sheet name; defaults to the active sheet
    :param energy_unit: unit of un-suffixed energy columns (``"wh"`` by default, or ``"kwh"``)
    """
    if kind not in METRICS:
        raise ValueError(f"unknown report kind {kind!r}; expected one of {sorted(METRICS)}")
    spec = METRICS[kind]

    title, rows = _read_rows(source, sheet)
    rows = list(rows)
    start = next((i for i, row in enumerate(rows) if not all(_is_blank(c) for c in row)), None)
    if start is None:
        raise ReportFormatError("the sheet is empty")

    header = rows[start]
    warnings: list[ParseWarning] = []
    date_columns = _select_date_columns(header, spec, energy_unit, warnings)
    if not date_columns:
        raise ReportFormatError("no date columns found in the header row")
    ordered_days = sorted(date_columns)

    date_indexes = {idx for idx, _, _ in date_columns.values()}
    field_columns: dict[int, str] = {}
    for index, cell in enumerate(header):
        if index in date_indexes:
            continue
        name = field_for_header(cell)
        if name and name not in field_columns.values():
            field_columns[index] = name

    parsed_rows: list[DeviceRow] = []
    seen_keys: dict[str, int] = {}

    for offset, row in enumerate(rows[start + 1 :], start=start + 2):  # 1-based sheet row numbers
        if all(_is_blank(c) for c in row):
            continue

        info = DeviceInfo()
        for index, name in field_columns.items():
            if index < len(row):
                setattr(info, name, _to_text(row[index]))
        if not info.key:
            warnings.append(ParseWarning(offset, "row has no serial number or device id; skipped"))
            continue
        if info.key in seen_keys:
            first_row = seen_keys[info.key]
            warnings.append(
                ParseWarning(offset, f"duplicate of row {first_row} ({info.key}); skipped")
            )
            continue
        seen_keys[info.key] = offset

        values: dict[date, float | None] = {}
        for day in ordered_days:
            index, factor, label = date_columns[day]
            cell = row[index] if index < len(row) else None
            number, problem = coerce_number(cell)
            if problem:
                warnings.append(ParseWarning(offset, problem, column=label))
            values[day] = None if number is None else round(number * factor, 6)

        parsed_rows.append(DeviceRow(info=info, values=values, row_number=offset))

    return Report(
        kind=kind,
        unit=spec.target_unit,
        dates=ordered_days,
        rows=parsed_rows,
        warnings=warnings,
        sheet=title,
    )
