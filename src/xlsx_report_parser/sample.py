"""Generate realistic-looking demo workbooks with entirely fictional devices.

The three files intentionally use different header styles (Russian / English / ISO dates /
unit suffixes) so the parser's flexibility can be seen on real files, not just in unit tests.
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

MODELS = ["Vendor X-100", "Vendor X-200 Pro", "Acme Hasher 7"]
CLIENTS = ["Demo Client A", "Demo Client B", "Demo Client C"]


def _device(index: int) -> dict:
    return {
        "id": f"DEV-{index:04d}",
        "serial": f"SN-DEMO-{index:04d}",
        "model": MODELS[index % len(MODELS)],
        "client": CLIENTS[index % len(CLIENTS)],
        "ip": f"10.20.{index // 250}.{index % 250 + 1}",
        "rack": f"R{index // 4 + 1}",
        "shelf": f"S{index % 4 + 1}",
    }


def generate_sample(
    out_dir: str | Path,
    *,
    start: date = date(2026, 3, 1),
    days: int = 7,
    devices: int = 6,
    seed: int = 42,
) -> list[Path]:
    """Write ``energy.xlsx``, ``uptime.xlsx`` and ``hashrate.xlsx``; return their paths."""
    rng = random.Random(seed)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    day_list = [start + timedelta(days=i) for i in range(days)]
    fleet = [_device(i + 1) for i in range(devices)]

    # --- energy: Russian headers, dd.mm.yyyy dates, values in watt-hours
    energy = Workbook()
    ws = energy.active
    ws.title = "Данные"
    ws.append(
        ["ID", "Серийный номер", "Модель", "Клиент", "Стойка", "Полка"]
        + [d.strftime("%d.%m.%Y") for d in day_list]
    )
    for dev in fleet:
        watts = rng.uniform(2800, 3400)  # average draw; a full day is watts * 24 Wh
        ws.append(
            [dev["id"], dev["serial"], dev["model"], dev["client"], dev["rack"], dev["shelf"]]
            + [round(watts * 24 * rng.uniform(0.9, 1.0)) for _ in day_list]
        )

    # --- uptime: English headers, ISO dates, one gap and one placeholder
    uptime = Workbook()
    ws = uptime.active
    ws.title = "Uptime"
    ws.append(["Device ID", "SN", "Model", "Client"] + [d.isoformat() for d in day_list])
    for n, dev in enumerate(fleet):
        row = [dev["id"], dev["serial"], dev["model"], dev["client"]]
        for i, _ in enumerate(day_list):
            if n == 1 and i == 2:
                row.append(None)  # telemetry gap
            elif n == 2 and i == 3:
                row.append("-")  # placeholder text
            else:
                row.append(round(rng.uniform(21.5, 24.0), 2))
        ws.append(row)

    # --- hashrate: unit in the header; a second unit column per day that must be ignored
    hashrate = Workbook()
    ws = hashrate.active
    ws.title = "Hashrate"
    header = ["ID", "SN", "M", "C"]
    for d in day_list:
        header += [f"{d.strftime('%d.%m.%Y')} (Th/s)", f"{d.strftime('%d.%m.%Y')} (Mh/s)"]
    ws.append(header)
    for dev in fleet:
        base = rng.uniform(90, 110)
        row = [dev["id"], dev["serial"], dev["model"], dev["client"]]
        for _ in day_list:
            value = round(base * rng.uniform(0.95, 1.02), 2)
            row += [value, round(value * 1_000_000)]
        ws.append(row)

    paths = []
    for name, workbook in (("energy", energy), ("uptime", uptime), ("hashrate", hashrate)):
        path = out / f"{name}.xlsx"
        workbook.save(path)
        paths.append(path)
    return paths
